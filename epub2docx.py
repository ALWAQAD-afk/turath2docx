from pathlib import Path
from zipfile import ZipFile
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import re
import sys
import warnings
from bs4 import XMLParsedAsHTMLWarning

warnings.filterwarnings("ignore", category=XMLParsedAsHTMLWarning)

EPUB = Path("6684.epub")
OUTPUT = Path("المنتقى-epub-50-v3.docx")

MAX_PAGES = 50

FONT = "Traditional Arabic"

# نجعلها مضبوطة مبدئيًا على A4.
PAGE_WIDTH = Mm(210)
PAGE_HEIGHT = Mm(297)

TOP_MARGIN = Mm(14)
BOTTOM_MARGIN = Mm(14)
RIGHT_MARGIN = Mm(16)
LEFT_MARGIN = Mm(16)

BODY_SIZE = 14
HEADING_SIZE = 17
FOOTER_SIZE = 9


def clean(text):
    if not text:
        return ""

    text = (
        text.replace("\u200f", "")
            .replace("\u200e", "")
            .replace("\ufeff", "")
            .replace("\xa0", " ")
    )

    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\s*\n\s*", " ", text)

    return text.strip()


def normalize(text):
    text = clean(text)
    text = re.sub(r"[ـ\s]+", "", text)
    return text


def set_rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()

    bidi = pPr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        pPr.append(bidi)

    bidi.set(qn("w:val"), "1")


def style_run(run, size, bold=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold

    rPr = run._r.get_or_add_rPr()

    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)

    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rFonts.set(qn(f"w:{attr}"), FONT)

    rtl = rPr.find(qn("w:rtl"))
    if rtl is None:
        rtl = OxmlElement("w:rtl")
        rPr.append(rtl)

    rtl.set(qn("w:val"), "1")


def configure_section(section):
    section.page_width = PAGE_WIDTH
    section.page_height = PAGE_HEIGHT

    section.top_margin = TOP_MARGIN
    section.bottom_margin = BOTTOM_MARGIN
    section.right_margin = RIGHT_MARGIN
    section.left_margin = LEFT_MARGIN

    section.header_distance = Mm(6)
    section.footer_distance = Mm(6)


def set_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.clear()

    set_rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = p.add_run("الجزء 1  |  الصفحة ")
    style_run(run, FOOTER_SIZE)

    run = p.add_run()
    style_run(run, FOOTER_SIZE)

    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "

    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")

    placeholder = OxmlElement("w:t")
    placeholder.text = "1"

    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(placeholder)
    run._r.append(fld_end)


def add_page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)

    run = p.add_run()
    run.add_break(WD_BREAK.PAGE)

def add_body(doc, text):
    text = clean(text)

    if not text:
        return

    p = doc.add_paragraph()
    set_rtl(p)

    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    fmt = p.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(2)
    fmt.line_spacing = 1.0

    run = p.add_run(text)
    style_run(run, BODY_SIZE)


def add_heading(doc, text, level=2):
    text = clean(text)

    if not text:
        return

    p = doc.add_paragraph()
    set_rtl(p)

    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    fmt = p.paragraph_format
    fmt.space_before = Pt(4)
    fmt.space_after = Pt(3)
    fmt.keep_with_next = True

    size = HEADING_SIZE

    if level == 1:
        size = 19
    elif level >= 3:
        size = 16

    run = p.add_run(text)
    style_run(run, size, bold=True)


def is_junk(text):
    t = clean(text)

    if not t:
        return True

    # رؤوس الصفحات المتكررة
    if normalize(t) in {
        "[المنتقى]",
        "المنتقى",
    }:
        return True

    # خط الفصل
    if re.fullmatch(r"ـ+", t):
        return True

    # نقطة منفردة
    if t == ".":
        return True

    # سطر نقاط فقط
    if re.fullmatch(r"(?:\.\s*){4,}", t):
        return True

    # عبارة نهاية الصفحة:
    # الجزء: 1 - الصفحة: 2
    if re.fullmatch(
        r"الجزء\s*:\s*\d+\s*-\s*الصفحة\s*:\s*\d+",
        t
    ):
        return True

    return False


def page_files(z):
    out = []

    for name in z.namelist():
        low = name.lower()

        if (
            "/text/page_" in low
            and low.endswith(".xhtml")
        ):
            m = re.search(
                r"page_(\d+)_(\d+)\.xhtml$",
                low
            )

            if m:
                vol = int(m.group(1))
                page = int(m.group(2))

                out.append(
                    (vol, page, name)
                )

    out.sort(key=lambda x: (x[0], x[1]))

    return out


def extract_blocks(raw):
    # XHTML = XML، فلا نستخدم HTML parser.
    soup = BeautifulSoup(raw, "xml")

    body = soup.find("body")

    if body is None:
        return []

    blocks = []

    # نتعامل فقط مع العناصر العليا داخل body،
    # فلا نكرر النص بسبب العناصر المتداخلة.
    for node in body.find_all(recursive=False):

        if not getattr(node, "name", None):
            continue

        tag = node.name.lower()

        text = clean(
            node.get_text(" ", strip=True)
        )

        if is_junk(text):
            continue

        if tag in {
            "h1", "h2", "h3",
            "h4", "h5", "h6"
        }:
            blocks.append({
                "type": "heading",
                "level": int(tag[1]),
                "text": text,
            })
            continue

        # span data-type="title" هو رأس الصفحة المتكرر.
        if (
            tag == "span"
            and node.get("data-type") == "title"
        ):
            continue

        if tag == "p":
            blocks.append({
                "type": "paragraph",
                "text": text,
            })

    return blocks


def remove_empty_first_paragraph(doc):
    if not doc.paragraphs:
        return

    p = doc.paragraphs[0]

    if not p.text.strip():
        p._element.getparent().remove(p._element)


def main():

    if not EPUB.exists():
        print(f"❌ لم أجد {EPUB}")
        sys.exit(1)

    doc = Document()

    # Section واحدة فقط للكتاب كله
    section = doc.sections[0]
    configure_section(section)
    set_footer(section)

    with ZipFile(EPUB, "r") as z:

        all_pages = page_files(z)

        print(f"✓ Source pages found: {len(all_pages)}")

        pages = all_pages[:MAX_PAGES]

        # فحص صارم قبل إنشاء Word
        keys = [(v, pg) for v, pg, _ in pages]
        names = [name for _, _, name in pages]

        if len(names) != len(set(names)):
            raise SystemExit(
                "❌ Duplicate XHTML source files detected"
            )

        if MAX_PAGES == 50:
            expected = [(1, n) for n in range(1, 51)]

            if keys != expected:
                print("❌ Source sequence is not ج1 ص1 → ص50")
                print("Actual:", keys)
                sys.exit(1)

        print("✓ Source sequence verified: ج1 ص1 → ص50")
        print("✓ No duplicate XHTML source files")
        print("✓ Using ONE Word section")
        print("✓ Using ordinary page breaks")

        for index, (volume, page, name) in enumerate(pages):

            raw = z.read(name)
            blocks = extract_blocks(raw)

            for block in blocks:

                if block["type"] == "heading":
                    add_heading(
                        doc,
                        block["text"],
                        block["level"]
                    )
                else:
                    add_body(
                        doc,
                        block["text"]
                    )

            print(
                f"[{index + 1:02}/{len(pages)}] "
                f"ج{volume} ص{page}  "
                f"blocks={len(blocks)}"
            )

            # لا نضيف فاصلًا بعد الصفحة الأخيرة
            if index < len(pages) - 1:
                add_page_break(doc)

    remove_empty_first_paragraph(doc)

    doc.save(OUTPUT)

    print()
    print("✓ DONE")
    print(f"✓ Source pages written: {len(pages)}")
    print("✓ Sections: 1")
    print(f"✓ Explicit page breaks: {len(pages) - 1}")
    print(f"✓ Output: {OUTPUT}")
    print()
    print(
        "المطلوب الآن من العارض: 50/50. "
        "إذا ظهر 51/51 بعد إزالة Section Breaks، "
        "فإحدى صفحات المصدر نفسها overflowed."
    )


if __name__ == "__main__":
    main()
