from pathlib import Path
from zipfile import ZipFile
from bs4 import BeautifulSoup
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import re
import sys

EPUB = Path("6684.epub")
OUTPUT = Path("المنتقى-epub-50-v1.docx")

# أول اختبار فقط
MAX_PAGES = 50

FONT = "Traditional Arabic"
BODY_SIZE = 15
HEADING_SIZE = 17
PAGE_MARKER_SIZE = 9


def natural_key(s):
    return [
        int(x) if x.isdigit() else x.lower()
        for x in re.split(r"(\d+)", s)
    ]


def clean(text):
    if not text:
        return ""

    text = text.replace("\u200f", "")
    text = text.replace("\u200e", "")
    text = text.replace("\ufeff", "")
    text = text.replace("\xa0", " ")

    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n+", "\n", text)

    return text.strip()


def normalize(text):
    text = clean(text)
    text = re.sub(r"[\[\]﴿﴾()（）]", "", text)
    text = re.sub(r"[ـ\s]+", "", text)
    return text


def set_rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()

    bidi = pPr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        pPr.append(bidi)

    bidi.set(qn("w:val"), "1")


def style_run(run, size, bold=False, italic=False):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic

    rPr = run._r.get_or_add_rPr()

    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)

    for attr in (
        "ascii",
        "hAnsi",
        "eastAsia",
        "cs"
    ):
        rFonts.set(qn(f"w:{attr}"), FONT)

    rtl = rPr.find(qn("w:rtl"))
    if rtl is None:
        rtl = OxmlElement("w:rtl")
        rPr.append(rtl)

    rtl.set(qn("w:val"), "1")


def add_body(doc, text):
    text = clean(text)

    if not text:
        return

    p = doc.add_paragraph()
    set_rtl(p)

    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15

    run = p.add_run(text)
    style_run(run, BODY_SIZE)


def add_heading(doc, text, level=2):
    text = clean(text)

    if not text:
        return

    p = doc.add_paragraph()
    set_rtl(p)

    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True

    size = HEADING_SIZE

    if level == 1:
        size += 3
    elif level >= 3:
        size -= 1

    run = p.add_run(text)
    style_run(run, size, bold=True)


def add_page_marker(doc, volume, page):
    p = doc.add_paragraph()
    set_rtl(p)

    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)

    run = p.add_run(f"﴿ج{volume} ص{page}﴾")
    style_run(
        run,
        PAGE_MARKER_SIZE,
        italic=True
    )


def is_running_header(text):
    return normalize(text) == "المنتقى"


def get_page_info(soup, filename, fallback_index):
    """
    نحاول أخذ الجزء/الصفحة من بيانات EPUB نفسه.
    وإن لم نجدها نستخدم fallback للاختبار فقط.
    """

    volume = None
    page = None

    # البحث في النص أو السمات عن معلومات الصفحة.
    candidates = []

    for tag in soup.find_all(True):
        for key, value in tag.attrs.items():
            if isinstance(value, list):
                value = " ".join(value)

            candidates.append(
                f"{key}={value}"
            )

    candidates.append(soup.get_text(" ", strip=True)[:500])

    blob = " ".join(candidates)

    patterns = [
        r"ج(?:زء)?\s*(\d+)\s*[^\d]{0,10}ص(?:فحة)?\s*(\d+)",
        r"vol(?:ume)?[_\-:= ]*(\d+).*?page[_\-:= ]*(\d+)",
    ]

    for pattern in patterns:
        m = re.search(
            pattern,
            blob,
            flags=re.I
        )

        if m:
            volume = int(m.group(1))
            page = int(m.group(2))
            break

    # محاولة من اسم الملف
    if page is None:
        nums = re.findall(r"\d+", filename)

        if nums:
            page = int(nums[-1])

    if volume is None:
        volume = 1

    if page is None:
        page = fallback_index

    return volume, page


def useful_xhtml_files(z):
    files = []

    for name in z.namelist():
        low = name.lower()

        # نريد صفحات متن الكتاب فقط:
        # OEBPS/text/page_1_00001.xhtml
        if (
            "/text/page_" in low
            and low.endswith(".xhtml")
        ):
            files.append(name)

    files.sort(key=natural_key)
    return files

def extract_blocks(soup):
    """
    نستعمل البنية الصريحة في EPUB:
      h1..h6 = عنوان
      p      = فقرة
      data-type=title = عنوان/رأس

    لا ندمج الفقرات بالتخمين.
    """

    blocks = []

    # نفضّل body إن وجد.
    root = soup.body or soup

    elements = root.find_all(
        ["h1", "h2", "h3", "h4", "h5", "h6", "p"],
        recursive=True
    )

    for el in elements:
        text = clean(
            el.get_text(" ", strip=True)
        )

        if not text:
            continue

        tag = el.name.lower()

        if tag.startswith("h"):
            level = int(tag[1])

            blocks.append({
                "type": "heading",
                "level": level,
                "text": text
            })

            continue

        # <p> قد يحتوي data-type=title داخله.
        title_node = el.find(
            attrs={"data-type": "title"}
        )

        if title_node:
            title_text = clean(
                title_node.get_text(
                    " ",
                    strip=True
                )
            )

            if title_text:
                blocks.append({
                    "type": "special-title",
                    "text": title_text
                })

            # إن كان هناك نص آخر حقيقي في نفس p
            # نحافظ عليه.
            clone = BeautifulSoup(
                str(el),
                "lxml"
            )

            for x in clone.find_all(
                attrs={"data-type": "title"}
            ):
                x.decompose()

            remainder = clean(
                clone.get_text(
                    " ",
                    strip=True
                )
            )

            if remainder:
                blocks.append({
                    "type": "paragraph",
                    "text": remainder
                })

            continue

        blocks.append({
            "type": "paragraph",
            "text": text
        })

    return blocks


def main():
    if not EPUB.exists():
        print(f"❌ لم أجد: {EPUB}")
        print()
        print(
            "ضع 6684.epub داخل ~/turath2docx"
        )
        sys.exit(1)

    doc = Document()

    # الهوامش
    section = doc.sections[0]
    section.top_margin = Pt(36)
    section.bottom_margin = Pt(36)
    section.left_margin = Pt(42)
    section.right_margin = Pt(42)

    # عنوان الكتاب
    p = doc.add_paragraph()
    set_rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    r = p.add_run("المنتقى شرح الموطإ")
    style_run(r, 24, bold=True)

    p = doc.add_paragraph()
    set_rtl(p)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    r = p.add_run("أبو الوليد الباجي")
    style_run(r, 17)

    seen_first_muntaqa = False

    with ZipFile(EPUB, "r") as z:
        files = useful_xhtml_files(z)

        print(
            f"✓ HTML/XHTML files found: {len(files)}"
        )

        processed = 0

        for name in files:
            raw = z.read(name)

            soup = BeautifulSoup(
                raw,
                "lxml"
            )

            blocks = extract_blocks(soup)

            if not blocks:
                continue

            # نستبعد ملفات الغلاف/الفهرس تلقائيًا
            # ما لم يظهر فيها نص كتاب حقيقي.
            total_text = " ".join(
                b["text"] for b in blocks
            )

            if len(total_text) < 80:
                continue

            if processed >= MAX_PAGES:
                break

            processed += 1

            volume, page = get_page_info(
                soup,
                name,
                processed
            )

            add_page_marker(
                doc,
                volume,
                page
            )

            for block in blocks:
                typ = block["type"]
                text = block["text"]

                if typ == "special-title":
                    if is_running_header(text):
                        # نبقي أول [المنتقى] فقط.
                        if seen_first_muntaqa:
                            continue

                        seen_first_muntaqa = True

                    add_heading(
                        doc,
                        text,
                        level=3
                    )

                elif typ == "heading":
                    add_heading(
                        doc,
                        text,
                        block["level"]
                    )

                elif typ == "paragraph":
                    add_body(
                        doc,
                        text
                    )

            print(
                f"[{processed:02}/{MAX_PAGES}] "
                f"{name}  ->  ج{volume} ص{page}"
            )

    if processed == 0:
        print("❌ لم أجد صفحات كتاب قابلة للتحويل.")
        sys.exit(1)

    doc.save(OUTPUT)

    print()
    print("✓ DONE")
    print(f"✓ Pages processed: {processed}")
    print(f"✓ Output: {OUTPUT}")
    print()
    print(
        "مهم: هذه نسخة اختبار. "
        "لا نطلق الكتاب كاملًا قبل فحصها."
    )


if __name__ == "__main__":
    main()
